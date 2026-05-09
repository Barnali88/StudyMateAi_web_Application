"""
services/extractor.py - Extract plain text from PDF, DOCX, and TXT files safely
"""
from pathlib import Path
from config import get_settings

settings = get_settings()


def _limit(text: str) -> str:
    return text[: settings.max_extract_chars].strip()


def extract_text(file_path: str | Path, file_type: str) -> str:
    file_path = Path(file_path)
    ft = file_type.lower().lstrip(".")
    if ft == "pdf":
        return _extract_pdf(file_path)
    if ft == "docx":
        return _extract_docx(file_path)
    if ft == "txt":
        return _extract_txt(file_path)
    raise ValueError(f"Unsupported file type: {file_type}")


def _extract_pdf(path: Path) -> str:
    import fitz

    text_parts: list[str] = []
    total = 0
    doc = fitz.open(str(path))
    try:
        for page in doc:
            page_text = page.get_text("text") or ""
            if not page_text.strip():
                continue
            remaining = settings.max_extract_chars - total
            if remaining <= 0:
                break
            text_parts.append(page_text[:remaining])
            total += min(len(page_text), remaining)
    finally:
        doc.close()
    return _limit("\n".join(text_parts))


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    total = 0

    def add_text(value: str) -> None:
        nonlocal total
        value = value.strip()
        if not value or total >= settings.max_extract_chars:
            return
        remaining = settings.max_extract_chars - total
        parts.append(value[:remaining])
        total += min(len(value), remaining)

    for p in doc.paragraphs:
        add_text(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                add_text(cell.text)
    return _limit("\n".join(parts))


def _extract_txt(path: Path) -> str:
    try:
        text = path.read_text(encoding="utf-8", errors="ignore")
    except UnicodeDecodeError:
        text = path.read_text(encoding="latin-1", errors="ignore")
    return _limit(text)
